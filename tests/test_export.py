import csv
from datetime import date
import io

from openpyxl import load_workbook
import pytest

from prospectsignal import (
    ATTRIBUTION,
    DEMO_ATTRIBUTION,
    FREDDO_COLUMNS,
    Store,
    demo_icp,
    freddo_csv_bytes,
    freddo_frame,
    load_checklist,
    xlsx_bytes,
)
from prospectsignal.storage import assert_no_person_columns


@pytest.fixture
def register_shortlist(bulk_files):
    store = Store()
    store.import_bulk(*bulk_files, downloaded_at=date(2026, 10, 1))
    store.set_meta("dataset", "register")
    store.set_meta("downloaded_at", "2026-10-01")
    store.add_to_shortlist(["100011379", "100022745", "100056852"], "test")
    return store


def parse_csv(payload: bytes) -> list[dict[str, str]]:
    return list(csv.DictReader(io.StringIO(payload.decode("utf-8"))))


def test_freddo_csv_header_is_first_line_and_exact(register_shortlist):
    payload = freddo_csv_bytes(register_shortlist.shortlist_frame())
    assert not payload.startswith(b"\xef\xbb\xbf")  # no BOM: a generic importer must see the header first
    first_line = payload.decode("utf-8").splitlines()[0]
    assert first_line == ",".join(f'"{column}"' for column in FREDDO_COLUMNS)
    assert_no_person_columns(FREDDO_COLUMNS)


def test_freddo_csv_keeps_orgnr_and_postcode_as_text(register_shortlist):
    payload = freddo_csv_bytes(register_shortlist.shortlist_frame())
    rows = {row["org_nr"]: row for row in parse_csv(payload)}
    oslo = rows["100011379"]
    assert oslo["region_postcode"] == "0150"
    assert '"0150"' in payload.decode("utf-8") and '"100011379"' in payload.decode("utf-8")
    assert oslo["organization_name"] == "FIXTURE ÆRLIG ØL OG BRØD AS"
    assert oslo["nace"] == "11.050"
    assert oslo["employee_band"] == "20-49"
    assert oslo["entity_kind"] == "hovedenhet"
    sub = rows["100056852"]
    assert sub["entity_kind"] == "underenhet" and sub["parent_org_nr"] == "100011379"


def test_attribution_and_download_date_in_every_row(register_shortlist):
    for row in parse_csv(freddo_csv_bytes(register_shortlist.shortlist_frame())):
        assert "Brønnøysundregistrene" in row["brreg_source"]
        assert "NLOD" in row["brreg_source"]
        assert row["brreg_refreshed_at"] == "2026-10-01"


def test_export_never_contains_contact_values(register_shortlist):
    from conftest import FAKE_CONTACT_VALUES

    text = freddo_csv_bytes(register_shortlist.shortlist_frame()).decode("utf-8")
    for value in FAKE_CONTACT_VALUES:
        assert value not in text
    enk = [row for row in parse_csv(text.encode()) if row["org_nr"] == "100022745"][0]
    assert enk["street_address"] == ""


def test_units_gone_from_register_are_not_exported(register_shortlist):
    register_shortlist.remove_unit("100022745")
    frame = freddo_frame(register_shortlist.shortlist_frame())
    assert "100022745" not in set(frame["org_nr"])
    assert len(frame) == 2


def test_formula_injection_is_neutralised(register_shortlist):
    register_shortlist._cursor().execute("UPDATE units SET name = '=HYPERLINK(\"x\")' WHERE org_nr = '100011379'")
    rows = parse_csv(freddo_csv_bytes(register_shortlist.shortlist_frame()))
    assert rows[0]["organization_name"].startswith("'=")


def test_xlsx_has_source_sheet_with_attribution(register_shortlist):
    payload = xlsx_bytes(
        register_shortlist.shortlist_frame(),
        meta=register_shortlist.meta(),
        checklist_rows=load_checklist().rows(),
        icp_lines=demo_icp().describe(),
        exported_on=date(2026, 10, 1),
    )
    book = load_workbook(io.BytesIO(payload))
    assert book.sheetnames == ["Shortlist", "Source & licence", "ICP", "Outreach checklist"]
    source = {row[0]: row[1] for row in book["Source & licence"].iter_rows(min_row=2, values_only=True)}
    assert source["Attribution"] == ATTRIBUTION
    assert source["Register downloaded"] == "2026-10-01"
    assert "not imply" in source["Consent"] or "never implies" in source["Consent"]
    sheet = book["Shortlist"]
    headers = [cell.value for cell in sheet[1]]
    org_column = headers.index("org_nr") + 1
    assert sheet.cell(row=2, column=org_column).value == "100011379"
    assert isinstance(sheet.cell(row=2, column=org_column).value, str)


def test_demo_export_is_labelled_as_fictional(demo_store):
    demo_store.add_to_shortlist(demo_store.query(demo_icp())["org_nr"].head(3).tolist())
    rows = parse_csv(freddo_csv_bytes(demo_store.shortlist_frame()))
    assert all(row["brreg_source"] == DEMO_ATTRIBUTION for row in rows)
    assert all(row["org_nr"].startswith("0") for row in rows)
    book = load_workbook(io.BytesIO(xlsx_bytes(demo_store.shortlist_frame(), meta=demo_store.meta())))
    source = {row[0]: row[1] for row in book["Source & licence"].iter_rows(min_row=2, values_only=True)}
    assert source["Attribution"] == DEMO_ATTRIBUTION
